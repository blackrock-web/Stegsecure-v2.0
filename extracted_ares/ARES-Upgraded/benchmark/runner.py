"""
Full multi-model × multi-image benchmark runner.
Produces real metrics only — never fabricates values.
"""
from __future__ import annotations
import json, platform, sys, time, traceback
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional
from PIL import Image
import numpy as np

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from benchmark.registry import instantiate_models
from benchmark.metrics import full_metric_bundle
from benchmark.scoring import composite_scores, load_config


def make_sample_covers(n: int = 5, size: int = 256, seed: int = 42) -> List[Image.Image]:
    rng = np.random.RandomState(seed)
    covers = []
    bases = [(90, 110, 140), (120, 80, 60), (60, 100, 80), (140, 100, 120), (80, 90, 110)]
    for i in range(n):
        base = bases[i % len(bases)]
        arr = np.array(Image.new("RGB", (size, size), color=base), dtype=np.int16)
        for scale, amp in [(1, 35), (4, 18), (16, 10)]:
            noise = rng.randint(-amp, amp + 1, (max(1, size // scale), max(1, size // scale), 3))
            noise = np.repeat(np.repeat(noise, scale, axis=0), scale, axis=1)[:size, :size]
            arr = arr + noise
        covers.append(Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8)))
    return covers


def environment_info() -> dict:
    import torch
    return {
        "python": sys.version,
        "platform": platform.platform(),
        "pytorch": torch.__version__,
        "cuda_available": torch.cuda.is_available(),
        "cuda_device": torch.cuda.get_device_name(0) if torch.cuda.is_available() else None,
        "timestamp": datetime.utcnow().isoformat() + "Z",
    }


def run_benchmark(
    images: Optional[List[Image.Image]] = None,
    secret: str = None,
    password: str = None,
    out_dir: Optional[Path] = None,
) -> Dict[str, Any]:
    cfg = load_config()
    secret = secret or cfg["secret_default"]
    password = password or cfg["password"]
    size = cfg.get("image_size", 256)

    if images is None:
        images = make_sample_covers(5, size, cfg.get("random_seed", 42))
    # resize all to same size for fairness
    images = [im.convert("RGB").resize((size, size), Image.BILINEAR) for im in images[:5]]

    run_id = datetime.now().strftime("run_%Y%m%d_%H%M%S")
    out = Path(out_dir or ROOT / "results" / run_id)
    out.mkdir(parents=True, exist_ok=True)
    (out / "images").mkdir(exist_ok=True)
    (out / "plots").mkdir(exist_ok=True)

    print("=" * 60)
    print("StegResearchBenchmark — Full Comparison")
    print("=" * 60)
    env = environment_info()
    print(f"Device: {'CUDA ' + str(env['cuda_device']) if env['cuda_available'] else 'CPU'}")
    (out / "environment.json").write_text(json.dumps(env, indent=2))
    (out / "config.json").write_text(json.dumps(cfg, indent=2))

    print("\nLoading models…")
    device = "cuda" if env["cuda_available"] else "cpu"
    models = instantiate_models(device=device)

    raw_rows = []
    image_tables = {}

    for img_idx, cover in enumerate(images):
        cover.save(out / "images" / f"cover_{img_idx+1}.png")
        print(f"\n--- Image {img_idx+1}/{len(images)} ---")
        image_tables[f"Image_{img_idx+1}"] = []

        for model in models:
            print(f"  {model.model_name} …", end=" ", flush=True)
            try:
                result = model.evaluate(cover, secret, password)
                if result.get("status") != "ok":
                    print(f"FAILED ({result.get('error', '')[:60]})")
                    row = {
                        "image": img_idx + 1,
                        "model_id": model.model_id,
                        "model_name": model.model_name,
                        "status": "FAILED",
                        "error": result.get("error"),
                    }
                    raw_rows.append(row)
                    image_tables[f"Image_{img_idx+1}"].append(row)
                    continue

                stego = result["stego"]
                stego.save(out / "images" / f"stego_img{img_idx+1}_{model.model_id}.png")

                def _dec(img, pw):
                    dr = model.decode(img, pw)
                    return dr.secret if dr.status == "ok" else ""

                metrics = full_metric_bundle(
                    cover, stego,
                    payload_bits=result["payload_bits"],
                    recovery_accuracy=result["secret_recovery_accuracy"],
                    total_time_ms=result["total_time_ms"],
                    decode_fn=_dec,
                    secret=secret,
                    password=password,
                )
                row = {
                    "image": img_idx + 1,
                    "model_id": model.model_id,
                    "model_name": model.model_name,
                    "status": "ok",
                    **metrics,
                    "payload_bits": result["payload_bits"],
                    "encode_time_ms": result["encode_time_ms"],
                    "decode_time_ms": result["decode_time_ms"],
                }
                raw_rows.append(row)
                image_tables[f"Image_{img_idx+1}"].append(row)
                print(f"PSNR={metrics['PSNR']:.2f} SSIM={metrics['SSIM']:.4f} Rec={metrics['Secret_Recovery_Accuracy']:.0f}%")
            except Exception as e:
                print(f"EXCEPTION {e}")
                traceback.print_exc()
                row = {
                    "image": img_idx + 1,
                    "model_id": model.model_id,
                    "model_name": model.model_name,
                    "status": "FAILED",
                    "error": str(e),
                }
                raw_rows.append(row)
                image_tables[f"Image_{img_idx+1}"].append(row)

    # Aggregate per model
    metric_keys = [
        "PSNR", "SSIM", "MS_SSIM", "LPIPS", "MSE",
        "Payload_Capacity_bpp", "Secret_Recovery_Accuracy", "Robustness",
        "Inference_Time_ms", "Steganalysis_Detectability",
    ]
    aggregates = []
    model_ids = []
    for model in models:
        mid = model.model_id
        model_ids.append(mid)
        rows = [r for r in raw_rows if r.get("model_id") == mid and r.get("status") == "ok"]
        if not rows:
            aggregates.append({
                "model_id": mid, "model_name": model.model_name, "status": "NO_VALID_RUNS",
                "n_success": 0, "n_total": len(images),
            })
            continue
        agg = {
            "model_id": mid,
            "model_name": model.model_name,
            "status": "ok",
            "n_success": len(rows),
            "n_total": len(images),
        }
        for k in metric_keys:
            vals = [float(r[k]) for r in rows if k in r]
            if vals:
                agg[k] = round(float(np.mean(vals)), 6)
                agg[f"{k}_std"] = round(float(np.std(vals)), 6)
                agg[f"{k}_min"] = round(float(np.min(vals)), 6)
                agg[f"{k}_max"] = round(float(np.max(vals)), 6)
        aggregates.append(agg)

    # Composite score on successful aggregates
    ok_aggs = [a for a in aggregates if a.get("status") == "ok"]
    ranked = composite_scores(ok_aggs, cfg) if ok_aggs else []

    best = ranked[0] if ranked else None
    if best and len(ranked) > 1:
        if abs(ranked[0]["overall_score"] - ranked[1]["overall_score"]) < cfg.get("tie_tolerance", 0.01):
            best_label = "Statistical tie / no clear separation"
        else:
            best_label = best["model_name"]
    else:
        best_label = best["model_name"] if best else "N/A"

    summary = {
        "run_id": run_id,
        "n_images": len(images),
        "n_models": len(models),
        "best_measured_model": best_label,
        "best_score": best["overall_score"] if best else None,
        "ranked": ranked,
        "aggregates": aggregates,
    }

    # Save JSON
    (out / "raw_metrics.json").write_text(json.dumps(raw_rows, indent=2, default=str))
    (out / "aggregate_metrics.json").write_text(json.dumps(aggregates, indent=2, default=str))
    (out / "summary.json").write_text(json.dumps(summary, indent=2, default=str))

    # CSV
    try:
        import csv
        with open(out / "raw_results.csv", "w", newline="") as f:
            if raw_rows:
                keys = list(raw_rows[0].keys())
                w = csv.DictWriter(f, fieldnames=keys, extrasaction="ignore")
                w.writeheader()
                w.writerows(raw_rows)
        with open(out / "aggregate_results.csv", "w", newline="") as f:
            if ranked:
                keys = list(ranked[0].keys())
                w = csv.DictWriter(f, fieldnames=keys, extrasaction="ignore")
                w.writeheader()
                w.writerows(ranked)
    except Exception as e:
        print("CSV export warning:", e)

    # Excel if openpyxl available
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment
        wb = openpyxl.Workbook()
        # Summary
        ws = wb.active
        ws.title = "Summary"
        ws.append(["StegResearchBenchmark Results"])
        ws.append(["Run", run_id])
        ws.append(["Best measured model under configured benchmark", best_label])
        ws.append(["Best score", best["overall_score"] if best else ""])
        ws.append([])
        ws.append(["Model", "Overall Score", "Rank", "Avg PSNR", "Avg SSIM", "Avg Recovery", "Avg Robustness"])
        for r in ranked:
            ws.append([
                r.get("model_name"), r.get("overall_score"), r.get("rank"),
                r.get("PSNR"), r.get("SSIM"), r.get("Secret_Recovery_Accuracy"), r.get("Robustness"),
            ])
        # Per image
        for name, rows in image_tables.items():
            ws2 = wb.create_sheet(name[:31])
            if rows:
                keys = [k for k in rows[0].keys() if k != "error"]
                ws2.append(keys)
                for r in rows:
                    ws2.append([r.get(k, "") for k in keys])
        # Aggregate
        ws3 = wb.create_sheet("Aggregate")
        if ranked:
            keys = list(ranked[0].keys())
            ws3.append(keys)
            for r in ranked:
                ws3.append([r.get(k, "") for k in keys])
        wb.save(out / "comparison_results.xlsx")
        print(f"Excel saved: {out / 'comparison_results.xlsx'}")
    except Exception as e:
        print("Excel export skipped:", e)

    print("\n" + "=" * 60)
    print("BEST MEASURED MODEL UNDER THIS BENCHMARK CONFIGURATION")
    print(f"  → {best_label}")
    if best:
        print(f"  Overall score: {best['overall_score']}")
    print(f"Results directory: {out}")
    print("=" * 60)

    return {"out_dir": str(out), "summary": summary, "raw": raw_rows, "ranked": ranked}


if __name__ == "__main__":
    run_benchmark()
